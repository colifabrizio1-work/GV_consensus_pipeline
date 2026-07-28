from __future__ import annotations

import argparse
import logging
import os
import re
from copy import copy
from datetime import date, datetime
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook
from openpyxl.formula.translate import Translator
from openpyxl.styles import PatternFill, Side
from openpyxl.utils import get_column_letter

from gv_config import load_config

MONTH_COLUMNS = list(range(2, 15))  # B:N, 13 rolling months: -6 .. +6
ROLLING_OFFSETS = list(range(-6, 7))
SALES_ROWS = {"LLLY": 5, "LLY": 6, "LY": 7, "CY": 8}
WEEKLY_ROWS = {"LLLY": 10, "LLY": 11, "LY": 12, "CY": 13}
MINIMO_ROWS = {"LLLY": 27, "LLY": 28, "LY": 29, "CY": 30}
MINIMO_PCT_ROWS = {"LLLY": 32, "LLY": 33, "LY": 34, "CY": 35}
YEAR_KEYS = ["LLLY", "LLY", "LY", "CY"]
SELF_EMAIL = "fabrizio.coli@luxottica.com"
QUARTER_MONTH_FILLS = {
    1: PatternFill("solid", fgColor="FFF2CC"),
    2: PatternFill("solid", fgColor="FFF2CC"),
    3: PatternFill("solid", fgColor="FFF2CC"),
    4: PatternFill("solid", fgColor="D9E2F3"),
    5: PatternFill("solid", fgColor="D9E2F3"),
    6: PatternFill("solid", fgColor="D9E2F3"),
    7: PatternFill("solid", fgColor="D9EAD3"),
    8: PatternFill("solid", fgColor="D9EAD3"),
    9: PatternFill("solid", fgColor="D9EAD3"),
    10: PatternFill("solid", fgColor="FCE4D6"),
    11: PatternFill("solid", fgColor="FCE4D6"),
    12: PatternFill("solid", fgColor="FCE4D6"),
}


def setup_logger(log_dir: str, dry_run: bool) -> logging.Logger:
    Path(log_dir).mkdir(parents=True, exist_ok=True)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    suffix = "_dry_run" if dry_run else ""
    log_path = Path(log_dir) / f"generate_consensus_frames_gv_{timestamp}{suffix}.log"

    logger = logging.getLogger("generate_consensus_frames_gv")
    logger.setLevel(logging.INFO)
    logger.handlers.clear()

    formatter = logging.Formatter("%(asctime)s | %(levelname)s | %(message)s")
    file_handler = logging.FileHandler(log_path, encoding="utf-8")
    file_handler.setFormatter(formatter)
    stream_handler = logging.StreamHandler()
    stream_handler.setFormatter(formatter)
    logger.addHandler(file_handler)
    logger.addHandler(stream_handler)
    logger.info("Log file: %s", log_path)
    logger.gv_log_path = log_path
    return logger


def shift_yyyymm(yyyymm: int, delta: int) -> int:
    year = yyyymm // 100
    month = yyyymm % 100
    idx = year * 12 + (month - 1) + delta
    return (idx // 12) * 100 + (idx % 12 + 1)


def parse_yyyymm(value: str | int) -> int:
    text = str(value).strip()
    if not re.fullmatch(r"\d{6}", text):
        raise ValueError(f"Invalid YYYYMM: {value!r}")
    month = int(text[-2:])
    if month < 1 or month > 12:
        raise ValueError(f"Invalid month in YYYYMM: {value!r}")
    return int(text)


def normalize_datest(series: pd.Series) -> pd.Series:
    return series.astype(str).str.strip().str.zfill(6)


def normalize_upc(series: pd.Series) -> pd.Series:
    return series.astype(str).str.strip()


def normalize_brand(series: pd.Series) -> pd.Series:
    return series.fillna("").astype(str).str.strip().str.casefold()


def current_fiscal_yyyymm(fiscal_calendar_path: str) -> int:
    today = pd.Timestamp(date.today())
    monthly = pd.read_excel(fiscal_calendar_path, sheet_name="Monthly")
    monthly["date start"] = pd.to_datetime(monthly["date start"])
    monthly["date end"] = pd.to_datetime(monthly["date end"])
    match = monthly[(monthly["date start"] <= today) & (monthly["date end"] >= today)]
    if match.empty:
        raise ValueError(f"No fiscal month found for today {today.date()} in {fiscal_calendar_path}")
    return int(match.iloc[0]["Year_month"])


def load_month_calendar(fiscal_calendar_path: str) -> tuple[pd.DataFrame, dict[int, dict[str, Any]], dict[str, int]]:
    weekly = pd.read_excel(fiscal_calendar_path, sheet_name="Weekly")
    monthly = pd.read_excel(fiscal_calendar_path, sheet_name="Monthly")
    weekly["Year F_Week"] = weekly["Year F_Week"].astype(str).str.strip()
    weekly["Year F_Month"] = weekly["Year F_Month"].astype(int)
    weekly["F_Week number"] = pd.to_numeric(weekly["F_Week number"], errors="coerce").astype("Int64")
    weekly["N settimane nel mese"] = pd.to_numeric(weekly["N settimane nel mese"], errors="coerce")

    monthly["Year_month"] = monthly["Year_month"].astype(int)
    month_info = {
        int(row["Year_month"]): {
            "name_it": str(row.get("Nome Mese", "")),
            "name_en": str(row.get("Month Name", "")),
            "weeks": int(row.get("N settimane nel mese", 0)) if "N settimane nel mese" in monthly.columns and pd.notna(row.get("N settimane nel mese")) else None,
        }
        for _, row in monthly.iterrows()
    }
    week_to_month = {str(row["Year F_Week"]): int(row["Year F_Month"]) for _, row in weekly.iterrows()}
    weeks_by_month = weekly.groupby("Year F_Month")["Year F_Week"].apply(lambda s: sorted(s.astype(str).tolist())).to_dict()
    for ym, weeks in weeks_by_month.items():
        month_info.setdefault(int(ym), {})["weeks_list"] = weeks
        month_info.setdefault(int(ym), {})["weeks"] = len(weeks)
    return weekly, month_info, week_to_month


def load_defill_pairs(cfg: dict) -> pd.DataFrame:
    info = cfg["manual_input"]["brand_defill"]
    df = pd.read_excel(info["path"], sheet_name=info["sheet"], dtype=str)
    df = df[[info["datest_column"], info["brand_column"]]].dropna()
    out = pd.DataFrame({
        "Client Datest": normalize_datest(df[info["datest_column"]]),
        "Brand_norm": normalize_brand(df[info["brand_column"]]),
    })
    return out.drop_duplicates()


def load_item_map(cfg: dict) -> pd.DataFrame:
    path = cfg["official_borrowed"]["anagrafica_base_parquet"]
    df = pd.read_parquet(path, columns=["UPC", "Brand", "Product Type"])
    df = df.dropna(subset=["UPC"]).copy()
    df["UPC"] = normalize_upc(df["UPC"])
    df["Brand"] = df["Brand"].fillna("").astype(str).str.strip()
    df["Product Type"] = df["Product Type"].fillna("unknown").astype(str).str.strip()
    df = df.drop_duplicates(subset=["UPC"], keep="last")

    wearable_path = cfg["official_borrowed"].get("brand_wearables")
    wearable_brands: set[str] = set()
    if wearable_path:
        wearable_df = pd.read_excel(wearable_path, dtype=str)
        first_col = wearable_df.columns[0]
        wearable_brands = set(normalize_brand(wearable_df[first_col]).dropna().tolist())
    df["Brand_norm"] = normalize_brand(df["Brand"])
    df["is_wearable"] = df["Brand_norm"].isin(wearable_brands)
    df = df.drop(columns=["Brand_norm"])
    return df


def enrich_and_filter_items(df: pd.DataFrame, item_map: pd.DataFrame, defill_pairs: pd.DataFrame, logger: logging.Logger, label: str) -> pd.DataFrame:
    before = len(df)
    out = df.merge(item_map, on="UPC", how="left")

    missing_mask = out["Brand"].isna()
    missing_rows = int(missing_mask.sum())
    missing_upcs = int(out.loc[missing_mask, "UPC"].nunique())
    logger.info("Anagrafica match %s: rows=%s missing_rows=%s missing_unique_upc=%s", label, before, missing_rows, missing_upcs)
    if missing_upcs:
        log_path = Path(getattr(logger, "gv_log_path", "")).parent if getattr(logger, "gv_log_path", None) else Path.cwd()
        report_path = log_path / f"missing_upc_{label}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.csv"
        out.loc[missing_mask, ["Client Datest", "UPC"]].drop_duplicates().sort_values(["Client Datest", "UPC"]).to_csv(report_path, index=False)
        logger.warning("Missing UPC report %s: %s", label, report_path)

    out["Product Type"] = out["Product Type"].fillna("unknown").astype(str).str.strip()
    accessory_mask = out["Product Type"].str.casefold().eq("accessories")
    wearable_mask = out["is_wearable"].fillna(False).astype(bool)
    accessory_rows = int(accessory_mask.sum())
    wearable_rows = int(wearable_mask.sum())
    out = out[~accessory_mask & ~wearable_mask].copy()
    logger.info(
        "Product exclusion %s: accessories_removed=%s wearables_removed=%s after_product_filter=%s",
        label,
        accessory_rows,
        wearable_rows,
        len(out),
    )

    out["Brand_norm"] = normalize_brand(out["Brand"])
    out = out.merge(defill_pairs.assign(_defill=1), on=["Client Datest", "Brand_norm"], how="left")
    defill_removed = int(out["_defill"].fillna(0).astype(bool).sum())
    out = out[out["_defill"].isna()].drop(columns=["Brand", "Product Type", "is_wearable", "Brand_norm", "_defill"])
    logger.info("Defill exclusion %s: before=%s removed=%s after=%s", label, before, defill_removed, len(out))
    return out


def load_sales(cfg: dict, item_map: pd.DataFrame, defill_pairs: pd.DataFrame, logger: logging.Logger) -> pd.DataFrame:
    sales_dir = Path(cfg["gv"]["parquet_sales"])
    parts = []
    for path in sorted(sales_dir.glob("GV_Sales_*.parquet")):
        part = pd.read_parquet(path, columns=["Client Datest", "UPC", "Fiscal Year", "Fiscal Month", "Fiscal Week", "Qty", "data_type"])
        parts.append(part)
    if not parts:
        raise FileNotFoundError(f"No GV sales parquet found in {sales_dir}")
    df = pd.concat(parts, ignore_index=True)
    df["Client Datest"] = normalize_datest(df["Client Datest"])
    df["UPC"] = normalize_upc(df["UPC"])
    invalid_upc = df["UPC"].str.lower().isin(["", "nan", "none", "nat"])
    if invalid_upc.any():
        logger.warning("Dropping sales rows with blank/invalid UPC: %s", int(invalid_upc.sum()))
        df = df[~invalid_upc].copy()
    df["Fiscal Year"] = pd.to_numeric(df["Fiscal Year"], errors="coerce").astype("Int64")
    df["Fiscal Month"] = pd.to_numeric(df["Fiscal Month"], errors="coerce").astype("Int64")
    df["Fiscal Month Num"] = (df["Fiscal Month"].astype("Int64") % 100).astype("Int64")
    df["Fiscal Week"] = df["Fiscal Week"].astype(str).str.strip()
    df["Qty"] = pd.to_numeric(df["Qty"], errors="coerce").fillna(0.0)
    df = enrich_and_filter_items(df, item_map, defill_pairs, logger, "sales")
    logger.info("Loaded sales after Defill: rows=%s", len(df))
    return df


def load_forecast(cfg: dict, item_map: pd.DataFrame, defill_pairs: pd.DataFrame, week_to_month: dict[str, int], logger: logging.Logger) -> pd.DataFrame:
    path = Path(cfg["parquet"]["forecast"])
    df = pd.read_parquet(path, columns=["Client Datest", "UPC", "Fiscal Week", "Qty", "data_type", "snapshot_timestamp"])
    df["Client Datest"] = normalize_datest(df["Client Datest"])
    df["UPC"] = normalize_upc(df["UPC"])
    invalid_upc = df["UPC"].str.lower().isin(["", "nan", "none", "nat"])
    if invalid_upc.any():
        logger.warning("Dropping forecast rows with blank/invalid UPC: %s", int(invalid_upc.sum()))
        df = df[~invalid_upc].copy()
    df["Fiscal Week"] = df["Fiscal Week"].astype(str).str.strip()
    df["Year_month"] = df["Fiscal Week"].map(week_to_month)
    missing = int(df["Year_month"].isna().sum())
    if missing:
        logger.warning("Forecast rows without fiscal week/month mapping: %s", missing)
        df = df[df["Year_month"].notna()].copy()
    df["Year_month"] = df["Year_month"].astype(int)
    df["Fiscal Year"] = df["Year_month"] // 100
    df["Fiscal Month Num"] = df["Year_month"] % 100
    df["Qty"] = pd.to_numeric(df["Qty"], errors="coerce").fillna(0.0)
    df = enrich_and_filter_items(df, item_map, defill_pairs, logger, "forecast")
    logger.info("Loaded forecast after Defill: rows=%s", len(df))
    return df


def build_lookups(sales: pd.DataFrame, forecast: pd.DataFrame) -> dict[str, dict[tuple, float]]:
    sales_month = (
        sales.groupby(["Client Datest", "data_type", "Fiscal Year", "Fiscal Month Num"], dropna=False)["Qty"].sum().to_dict()
    )
    sales_week = (
        sales.groupby(["Client Datest", "data_type", "Fiscal Year", "Fiscal Month Num", "Fiscal Week"], dropna=False)["Qty"].sum().to_dict()
    )
    forecast_month = (
        forecast.groupby(["Client Datest", "Year_month"], dropna=False)["Qty"].sum().to_dict()
    )
    return {"sales_month": sales_month, "sales_week": sales_week, "forecast_month": forecast_month}


def month_sequence(target_yyyymm: int) -> list[int]:
    return [shift_yyyymm(target_yyyymm, delta) for delta in ROLLING_OFFSETS]


def bucket_for_ym(ym: int, target_yyyymm: int) -> str:
    if ym == target_yyyymm:
        return "CM"
    ym_year = ym // 100
    target_year = target_yyyymm // 100
    if ym_year < target_year:
        return "LY"
    if ym_year == target_year:
        return "CY"
    return f"CY+{ym_year - target_year}"


def previous_consensus_path(output_path: Path, target_yyyymm: int) -> Path | None:
    prev_ym = shift_yyyymm(target_yyyymm, -1)
    path = output_path.with_name(f"Consensus_Frames_GV_{prev_ym}.xlsx")
    return path if path.exists() else None


def load_previous_values(prev_path: Path | None, datest: str) -> dict[int, dict[int, Any]]:
    if prev_path is None:
        return {}
    wb = load_workbook(prev_path, data_only=True, read_only=True)
    if datest not in wb.sheetnames:
        wb.close()
        return {}
    ws = wb[datest]
    out: dict[int, dict[int, Any]] = {}
    for col in range(2, ws.max_column + 1):
        ym = ws.cell(3, col).value
        if ym is None:
            continue
        try:
            ym_int = int(ym)
        except Exception:
            continue
        def _num(value: Any) -> float | None:
            try:
                if value in [None, ""]:
                    return None
                return float(value)
            except Exception:
                return None

        is_new_layout = "Weekly" in str(ws.cell(23, 1).value or "")
        sales_volume_row = 24 if is_new_layout else 23
        min_volume_row = 42 if is_new_layout else 39
        min_forecast_value_row = 38
        min_cy_row = 30 if is_new_layout else 29
        comment_row = 45 if is_new_layout else 42

        row24_value = ws.cell(sales_volume_row, col).value
        row19_value = _num(ws.cell(19, col).value)
        row8_value = _num(ws.cell(8, col).value)
        if row24_value in [None, ""] and row19_value is not None and row8_value not in [None, 0]:
            row24_value = row19_value / row8_value - 1

        row42_value = ws.cell(min_volume_row, col).value
        row38_value = _num(ws.cell(min_forecast_value_row, col).value)
        row30_value = _num(ws.cell(min_cy_row, col).value)
        if row42_value in [None, ""] and row38_value is not None and row30_value not in [None, 0]:
            row42_value = row38_value / row30_value - 1

        out[ym_int] = {
            21: ws.cell(21, col).value,
            22: ws.cell(22, col).value,
            24: row24_value,
            40: ws.cell(40, col).value if is_new_layout else ws.cell(36, col).value,
            41: ws.cell(41, col).value if is_new_layout else ws.cell(37, col).value,
            42: row42_value,
            45: ws.cell(comment_row, col).value,
        }
    wb.close()
    return out


def set_sheet_headers(ws, datest: str, target_yyyymm: int, months: list[int], month_info: dict[int, dict[str, Any]], year_map: dict[str, int]) -> None:
    ws.cell(1, 1).value = f"Consensus Frames GV - {datest} (wearable and defill excluded)"
    ws.row_dimensions[2].hidden = True
    ws.row_dimensions[3].hidden = True
    for key, row in SALES_ROWS.items():
        ws.cell(row, 1).value = f"Sales {year_map[key]}"
    for key, row in WEEKLY_ROWS.items():
        ws.cell(row, 1).value = f"Weekly Sales {year_map[key]}"
    for key, row in MINIMO_ROWS.items():
        ws.cell(row, 1).value = f"Sales da minimo {year_map[key]}"
    for key, row in MINIMO_PCT_ROWS.items():
        label = "Salesda minimo" if key == "LLY" else "Sales da minimo"
        ws.cell(row, 1).value = f"{label} {year_map[key]} (%)"
    ws.cell(15, 1).value = f"Sales {year_map['LY']} vs sales {year_map['LLY']}"
    ws.cell(16, 1).value = f"Sales {year_map['CY']} vs sales {year_map['LY']}"
    ws.cell(21, 1).value = f"Forecast Sales Tot   (% VS {year_map['LY']})"

    for col, ym in zip(MONTH_COLUMNS, months):
        month_name = month_info.get(ym, {}).get("name_it") or f"M{ym % 100}"
        month_num = ym % 100
        ws.cell(2, col).value = bucket_for_ym(ym, target_yyyymm)
        ws.cell(3, col).value = ym
        ws.cell(4, col).value = month_name
        ws.cell(4, col).fill = copy(QUARTER_MONTH_FILLS[month_num])
    for col in range(15, ws.max_column + 1):
        for row in [2, 3, 4]:
            ws.cell(row, col).value = None


def clear_month_area(ws) -> None:
    for row in range(5, 46):
        for col in MONTH_COLUMNS:
            ws.cell(row, col).value = None


def apply_special_vertical_separators(ws, months: list[int], target_yyyymm: int) -> None:
    thin_side = Side(style="thin", color="000000")
    quarter_side = Side(style="medium", color="000000")
    year_side = Side(style="double", color="000000")
    current_side = Side(style="mediumDashDot", color="EDB913")
    separator_styles = {"medium", "double", "mediumDashDot"}
    quarter_blank_rows = {9, 14, 17, 20, 26, 31, 36, 39, 44}

    # Quarter separator = right border after Mar/Jun/Sep/Dec.
    quarter_cols = {MONTH_COLUMNS[idx] for idx, ym in enumerate(months) if ym % 100 in (3, 6, 9, 12)}
    year_cols = {
        MONTH_COLUMNS[idx]
        for idx in range(1, len(months))
        if months[idx] // 100 != months[idx - 1] // 100
    }
    current_col = MONTH_COLUMNS[months.index(target_yyyymm)] if target_yyyymm in months else None

    # Clean stale thick separators from the template/previous layout before applying the current ones.
    for row in range(4, 46):
        for col in MONTH_COLUMNS:
            border = copy(ws.cell(row, col).border)
            if col != MONTH_COLUMNS[0] and border.left.style in separator_styles:
                border.left = thin_side
            if border.right.style in separator_styles:
                border.right = thin_side
            ws.cell(row, col).border = border

    # Priorita visiva: quarter nero, poi anno doppio, poi current month giallo.
    for col in quarter_cols:
        for row in range(4, 46):
            if row in quarter_blank_rows:
                continue
            border = copy(ws.cell(row, col).border)
            border.right = quarter_side
            ws.cell(row, col).border = border

    for col in year_cols:
        prev_col = col - 1
        for row in range(4, 46):
            if prev_col in MONTH_COLUMNS:
                border_prev = copy(ws.cell(row, prev_col).border)
                border_prev.right = Side(style=None)
                ws.cell(row, prev_col).border = border_prev
            border = copy(ws.cell(row, col).border)
            border.left = year_side
            ws.cell(row, col).border = border

    if current_col is not None:
        for row in range(4, 46):
            border = copy(ws.cell(row, current_col).border)
            border.left = current_side
            border.right = current_side
            ws.cell(row, current_col).border = border


def lookup_sales(lookups: dict[str, dict[tuple, float]], datest: str, data_type: str, fiscal_year: int, month_num: int) -> float:
    return float(lookups["sales_month"].get((datest, data_type, fiscal_year, month_num), 0.0))


def weekly_average(
    lookups: dict[str, dict[tuple, float]],
    datest: str,
    data_type: str,
    fiscal_year: int,
    month_num: int,
    ym: int,
    month_info: dict[int, dict[str, Any]],
) -> float:
    total = lookup_sales(lookups, datest, data_type, fiscal_year, month_num)
    weeks = month_info.get(ym, {}).get("weeks") or 0
    return total / weeks if weeks else 0.0


def sales_week_qty(
    lookups: dict[str, dict[tuple, float]],
    datest: str,
    data_type: str,
    fiscal_year: int,
    month_num: int,
    fiscal_week: str,
) -> float:
    return float(lookups["sales_week"].get((datest, data_type, fiscal_year, month_num, fiscal_week), 0.0))


def current_month_projection(
    lookups: dict[str, dict[tuple, float]],
    datest: str,
    target_yyyymm: int,
    ly_yyyymm: int,
    month_info: dict[int, dict[str, Any]],
    data_type: str = "Sales_hist",
) -> tuple[float, float]:
    cy_year = target_yyyymm // 100
    cy_month = target_yyyymm % 100
    ly_year = ly_yyyymm // 100
    ly_month = ly_yyyymm % 100
    cy_weeks = list(month_info.get(target_yyyymm, {}).get("weeks_list") or [])
    weeks = month_info.get(target_yyyymm, {}).get("weeks") or len(cy_weeks)
    ly_total = lookup_sales(lookups, datest, data_type, ly_year, ly_month)

    if not cy_weeks or not weeks:
        return ly_total, ly_total / weeks if weeks else 0.0

    closed_weeks = [
        week
        for week in cy_weeks
        if sales_week_qty(lookups, datest, data_type, cy_year, cy_month, week) != 0
    ]
    if not closed_weeks:
        return ly_total, ly_total / weeks if weeks else 0.0

    closed_cy = sum(
        sales_week_qty(lookups, datest, data_type, cy_year, cy_month, week)
        for week in closed_weeks
    )
    projected_total = closed_cy / len(closed_weeks) * weeks
    return projected_total, projected_total / weeks if weeks else 0.0


def write_values_and_formulas(
    ws,
    datest: str,
    target_yyyymm: int,
    months: list[int],
    year_map: dict[str, int],
    month_info: dict[int, dict[str, Any]],
    lookups: dict[str, dict[tuple, float]],
    previous_values: dict[int, dict[int, Any]],
) -> None:
    target_month = target_yyyymm % 100

    for col, ym in zip(MONTH_COLUMNS, months):
        month_num = ym % 100
        for key in YEAR_KEYS:
            fy = year_map[key]
            if key == "CY" and ym > target_yyyymm:
                # Future rolling months can repeat Jan/Feb of the following YYYYMM.
                # Keep the CY historical month populated so formulas can compare CY+1 vs CY.
                ws.cell(SALES_ROWS[key], col).value = lookup_sales(lookups, datest, "Sales_hist", fy, month_num)
                ws.cell(WEEKLY_ROWS[key], col).value = weekly_average(lookups, datest, "Sales_hist", fy, month_num, ym, month_info)
                ws.cell(MINIMO_ROWS[key], col).value = lookup_sales(lookups, datest, "Sales_minimo", fy, month_num)
            elif key == "CY" and ym == target_yyyymm:
                ly_same_month = int(f"{year_map['LY']}{target_month:02d}")
                projected_sales, projected_weekly = current_month_projection(
                    lookups,
                    datest,
                    target_yyyymm,
                    ly_same_month,
                    month_info,
                    "Sales_hist",
                )
                projected_minimo, _ = current_month_projection(
                    lookups,
                    datest,
                    target_yyyymm,
                    ly_same_month,
                    month_info,
                    "Sales_minimo",
                )
                ws.cell(SALES_ROWS[key], col).value = projected_sales
                ws.cell(WEEKLY_ROWS[key], col).value = projected_weekly
                ws.cell(MINIMO_ROWS[key], col).value = projected_minimo
            else:
                ws.cell(SALES_ROWS[key], col).value = lookup_sales(lookups, datest, "Sales_hist", fy, month_num)
                ws.cell(WEEKLY_ROWS[key], col).value = weekly_average(lookups, datest, "Sales_hist", fy, month_num, ym, month_info)
                ws.cell(MINIMO_ROWS[key], col).value = lookup_sales(lookups, datest, "Sales_minimo", fy, month_num)

    for col, ym in zip(MONTH_COLUMNS, months):
        letter = get_column_letter(col)
        ws.cell(15, col).value = f'=IFERROR(IF({letter}7/{letter}6-1=-1,"",{letter}7/{letter}6-1),"")'
        ws.cell(16, col).value = f'=IFERROR(IF({letter}8/{letter}7-1=-1,"",{letter}8/{letter}7-1),"")'
        ws.cell(32, col).value = f'=IFERROR({letter}27/{letter}5,"")'
        ws.cell(33, col).value = f'=IFERROR({letter}28/{letter}6,"")'
        ws.cell(34, col).value = f'=IFERROR({letter}29/{letter}7,"")'
        ws.cell(35, col).value = f'=IFERROR({letter}30/{letter}8,"")'

    previous_month = shift_yyyymm(target_yyyymm, -1)
    target_col = MONTH_COLUMNS[months.index(target_yyyymm)]
    previous_col = MONTH_COLUMNS[months.index(previous_month)] if previous_month in months else None
    sales_row_by_year = {year: SALES_ROWS[key] for key, year in year_map.items()}

    for col, ym in zip(MONTH_COLUMNS, months):
        letter = get_column_letter(col)
        prev = previous_values.get(ym, {})
        ws.cell(18, col).value = prev.get(21)
        ws.cell(19, col).value = prev.get(22)
        ws.cell(45, col).value = prev.get(45)
        forecast_qty = float(lookups["forecast_month"].get((datest, ym), 0.0))
        weeks_in_month = month_info.get(ym, {}).get("weeks") or len(month_info.get(ym, {}).get("weeks_list") or []) or 1
        if previous_col is not None and col < previous_col:
            ws.cell(21, col).value = None
            ws.cell(22, col).value = None
            ws.cell(23, col).value = None
            ws.cell(24, col).value = prev.get(24)
            ws.cell(25, col).value = None
        elif previous_col is not None and col in [previous_col, target_col]:
            ws.cell(21, col).value = f'={letter}16'
            ws.cell(22, col).value = f'={letter}8'
            ws.cell(23, col).value = f'=IFERROR({letter}22/{weeks_in_month},"")'
            ws.cell(24, col).value = f'=IFERROR({letter}19/{letter}8-1,"")'
            ws.cell(25, col).value = None if col == previous_col else f'=IFERROR({letter}22-{letter}19,"")'
        else:
            denom_row = sales_row_by_year.get((ym // 100) - 1, SALES_ROWS["LY"])
            ws.cell(21, col).value = f'=IFERROR(SUM({letter}22)/SUM({letter}{denom_row})-1,"")'
            if col < MONTH_COLUMNS[-1]:
                ws.cell(22, col).value = f'={letter}19'
            else:
                ws.cell(22, col).value = None
            ws.cell(23, col).value = f'=IFERROR({letter}22/{weeks_in_month},"")'
            ws.cell(24, col).value = None
            ws.cell(25, col).value = f'=IFERROR({letter}22-{letter}19,"")'

        if previous_col is not None and col < previous_col:
            ws.cell(37, col).value = None
            ws.cell(38, col).value = None
            ws.cell(40, col).value = None
            ws.cell(41, col).value = None
            ws.cell(42, col).value = prev.get(42)
            ws.cell(43, col).value = None
        elif previous_col is not None and col in [previous_col, target_col]:
            ws.cell(37, col).value = f'=IFERROR({letter}38/{letter}19,"")'
            ws.cell(38, col).value = forecast_qty
            ws.cell(40, col).value = f'={letter}35'
            ws.cell(41, col).value = f'={letter}30'
            ws.cell(42, col).value = f'=IFERROR({letter}38/{letter}30-1,"")'
            ws.cell(43, col).value = f'=IFERROR({letter}41-{letter}38,"")'
        else:
            ws.cell(37, col).value = f'=IFERROR({letter}38/{letter}19,"")'
            ws.cell(38, col).value = forecast_qty
            if col < MONTH_COLUMNS[-1]:
                ws.cell(40, col).value = f'={letter}37'
            else:
                ws.cell(40, col).value = None
            ws.cell(41, col).value = f'=IFERROR({letter}22*{letter}40,"")'
            ws.cell(42, col).value = None
            ws.cell(43, col).value = f'=IFERROR({letter}41-{letter}38,"")'

    for row in [15, 16, 18, 21, 24, 32, 33, 34, 35, 37, 40, 42]:
        for col in MONTH_COLUMNS:
            ws.cell(row, col).number_format = "0%"
    for row in [5, 6, 7, 8, 10, 11, 12, 13, 19, 22, 23, 25, 27, 28, 29, 30, 38, 41, 43]:
        for col in MONTH_COLUMNS:
            ws.cell(row, col).number_format = '#,##0'




def clone_worksheet_between_workbooks(source_ws, target_wb, title: str):
    ws = target_wb.create_sheet(title=title, index=0)
    for row in source_ws.iter_rows():
        for source_cell in row:
            target_cell = ws.cell(source_cell.row, source_cell.column)
            target_cell.value = source_cell.value
            if source_cell.has_style:
                target_cell._style = copy(source_cell._style)
            target_cell.number_format = source_cell.number_format
            target_cell.font = copy(source_cell.font)
            target_cell.fill = copy(source_cell.fill)
            target_cell.border = copy(source_cell.border)
            target_cell.alignment = copy(source_cell.alignment)
            target_cell.protection = copy(source_cell.protection)
            if source_cell.comment:
                target_cell.comment = copy(source_cell.comment)
    for merged_range in source_ws.merged_cells.ranges:
        ws.merge_cells(str(merged_range))
    for idx, dim in source_ws.column_dimensions.items():
        target_dim = ws.column_dimensions[idx]
        target_dim.width = dim.width
        target_dim.hidden = dim.hidden
        target_dim.outlineLevel = dim.outlineLevel
        target_dim.collapsed = dim.collapsed
    for idx, dim in source_ws.row_dimensions.items():
        target_dim = ws.row_dimensions[idx]
        target_dim.height = dim.height
        target_dim.hidden = dim.hidden
        target_dim.outlineLevel = dim.outlineLevel
        target_dim.collapsed = dim.collapsed
    ws.sheet_view.showGridLines = source_ws.sheet_view.showGridLines
    if source_ws.freeze_panes:
        ws.freeze_panes = source_ws.freeze_panes
    return ws


def month_names_for_completion(target_yyyymm: int, month_info: dict[int, dict[str, Any]]) -> list[str]:
    months = [shift_yyyymm(target_yyyymm, offset) for offset in (4, 5, 6)]
    return [month_info.get(ym, {}).get("name_it") or f"M{ym % 100}" for ym in months]


def is_completion_month_header(ws, row: int) -> bool:
    values = [ws.cell(row, col).value for col in range(2, 5)]
    non_empty = [str(value).strip() for value in values if value not in [None, ""]]
    if len(non_empty) < 2:
        return False
    month_names = {
        "gennaio", "febbraio", "marzo", "aprile", "maggio", "giugno",
        "luglio", "agosto", "settembre", "ottobre", "novembre", "dicembre",
    }
    return all(value.casefold() in month_names for value in non_empty)


def add_completion_sheet_from_previous(wb, prev_path: Path | None, target_yyyymm: int, month_info: dict[int, dict[str, Any]], logger: logging.Logger) -> None:
    if prev_path is None:
        logger.info("No previous completion sheet available")
        return
    prev_wb = load_workbook(prev_path)
    source_ws = None
    for sheet_name in prev_wb.sheetnames:
        if sheet_name.casefold() == "completamento":
            source_ws = prev_wb[sheet_name]
            break
    if source_ws is None:
        prev_wb.close()
        logger.info("Previous workbook has no completion sheet: %s", prev_path)
        return

    ws = clone_worksheet_between_workbooks(source_ws, wb, source_ws.title)
    prev_wb.close()

    ws.delete_cols(2, 1)
    for row in range(1, ws.max_row + 1):
        for col in (2, 3):
            cell = ws.cell(row, col)
            if isinstance(cell.value, str) and cell.value.startswith("="):
                old_col = get_column_letter(col + 1)
                new_col = get_column_letter(col)
                cell.value = Translator(cell.value, origin=f"{old_col}{row}").translate_formula(f"{new_col}{row}")
    ws.insert_cols(4, 1)
    for row in range(1, ws.max_row + 1):
        source = ws.cell(row, 3)
        target = ws.cell(row, 4)
        if source.has_style:
            target._style = copy(source._style)
        target.number_format = source.number_format
        target.font = copy(source.font)
        target.fill = copy(source.fill)
        target.border = copy(source.border)
        target.alignment = copy(source.alignment)
        target.protection = copy(source.protection)
        if isinstance(source.value, str) and source.value.startswith("="):
            target.value = Translator(source.value, origin=f"C{row}").translate_formula(f"D{row}")
        else:
            target.value = None

    month_names = month_names_for_completion(target_yyyymm, month_info)
    header_rows = [row for row in range(1, ws.max_row + 1) if is_completion_month_header(ws, row)]
    for row in header_rows:
        for col, month_name in zip(range(2, 5), month_names):
            ws.cell(row, col).value = month_name
    logger.info("Completion sheet rolled from %s | header_rows=%s | months=%s", prev_path, header_rows, month_names)

def copy_template_sheet(wb, template_ws, title: str):
    ws = wb.copy_worksheet(template_ws)
    ws.title = title
    return ws


def build_workbook(cfg: dict, yyyymm: int, output_path: Path, dry_run: bool, logger: logging.Logger) -> dict[str, Any]:
    fiscal_calendar = cfg["official_borrowed"]["fiscal_calendar"]
    _, month_info, week_to_month = load_month_calendar(fiscal_calendar)
    defill_pairs = load_defill_pairs(cfg)
    item_map = load_item_map(cfg)
    sales = load_sales(cfg, item_map, defill_pairs, logger)
    forecast = load_forecast(cfg, item_map, defill_pairs, week_to_month, logger)
    lookups = build_lookups(sales, forecast)

    datests = cfg["business_rules"]["allowed_gv_datests"]
    datests = sorted(str(d).zfill(6) for d in datests)
    months = month_sequence(yyyymm)
    cy = yyyymm // 100
    year_map = {"LLLY": cy - 3, "LLY": cy - 2, "LY": cy - 1, "CY": cy}

    template_path = Path(cfg["manual_input"]["consensus_template"]["path"])
    wb = load_workbook(template_path)
    template_ws = wb[cfg["manual_input"]["consensus_template"].get("sheet", "Consensus Frames")]
    prev_path = previous_consensus_path(output_path, yyyymm)
    if prev_path:
        logger.info("Previous consensus found for last approval: %s", prev_path)
    else:
        logger.info("No previous consensus found for last approval")

    add_completion_sheet_from_previous(wb, prev_path, yyyymm, month_info, logger)

    created = []
    for datest in datests:
        ws = copy_template_sheet(wb, template_ws, datest)
        clear_month_area(ws)
        set_sheet_headers(ws, datest, yyyymm, months, month_info, year_map)
        previous_values = load_previous_values(prev_path, datest)
        write_values_and_formulas(ws, datest, yyyymm, months, year_map, month_info, lookups, previous_values)
        apply_special_vertical_separators(ws, months, yyyymm)
        created.append(datest)

    wb.remove(template_ws)
    wb.active = 0
    output_path.parent.mkdir(parents=True, exist_ok=True)
    if not dry_run:
        tmp_path = output_path.with_name(output_path.name + ".tmp")
        wb.save(tmp_path)
        os.replace(tmp_path, output_path)
    wb.close()

    return {
        "output": str(output_path),
        "dry_run": dry_run,
        "datests": created,
        "months": months,
        "sales_rows_after_defill": len(sales),
        "forecast_rows_after_defill": len(forecast),
        "previous": str(prev_path) if prev_path else None,
    }


def main() -> int:
    parser = argparse.ArgumentParser(description="Generate GV consensus frames workbook.")
    parser.add_argument("--yyyymm", help="Consensus month in YYYYMM format. Defaults to current fiscal month.")
    parser.add_argument("--dry-run", action="store_true", help="Validate and build in memory without writing workbook.")
    args = parser.parse_args()

    cfg = load_config()
    logger = setup_logger(cfg["gv"]["log"], args.dry_run)
    try:
        target_yyyymm = parse_yyyymm(args.yyyymm) if args.yyyymm else current_fiscal_yyyymm(cfg["official_borrowed"]["fiscal_calendar"])
        output = Path(cfg["outputs"]["consensus_pattern"].format(yyyymm=target_yyyymm))
        logger.info("Starting GV consensus generation | yyyymm=%s | dry_run=%s", target_yyyymm, args.dry_run)
        result = build_workbook(cfg, target_yyyymm, output, args.dry_run, logger)
        logger.info("Completed GV consensus generation: %s", result)
        return 0
    except Exception:
        logger.exception("GV consensus generation failed")
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
